import csv
import io
import re
import logging
from celery import shared_task
from django.db import transaction
from .models import Questionnaire, Question

logger = logging.getLogger(__name__)

def parse_questionnaire(questionnaire):
    name = questionnaire.original_filename.lower()
    with questionnaire.source_file.open("rb") as source: data = source.read()
    rows = []
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() if value is not None else "" for value in row]
                    if any(values): rows.append(values)
        finally: workbook.close()
    elif name.endswith(".csv"):
        rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))) if any(cell.strip() for cell in row)]
    else:
        if name.endswith(".pdf"):
            import fitz
            with fitz.open(stream=data, filetype="pdf") as pdf: lines = "\n".join(page.get_text("text") for page in pdf).splitlines()
        elif name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs]
            lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for table in doc.tables for row in table.rows)
        else: lines = data.decode("utf-8-sig", errors="replace").splitlines()
        for line in lines:
            text = re.sub(r"^\s*(?:(?:question\s*)?\d+[.)\]:-]?\s+|[-*•]\s*)", "", line).strip()
            if len(text) >= 12 and ("?" in text or re.match(r"^(do|does|did|is|are|can|will|how|what|when|where|which|describe|explain|provide|list|identify|please)\b", text, re.I)):
                rows.append([line.strip().split(" ", 1)[0].strip(".):-") if line.strip()[:1].isdigit() else "", text])
    if not rows: return []
    header = [cell.lower().strip() for cell in rows[0]]
    question_col = next((i for i, value in enumerate(header) if any(key in value for key in ("question", "prompt", "requirement"))), None)
    number_col = next((i for i, value in enumerate(header) if any(key in value for key in ("number", "id", "ref"))), None)
    start = 1 if question_col is not None else 0
    question_col = question_col if question_col is not None else 0
    questions = []
    seen = set()
    for row in rows[start:]:
        if question_col >= len(row): continue
        text = row[question_col].strip()
        if len(text) < 8 or text.lower() in {"question", "prompt", "requirement"}: continue
        if not ("?" in text or len(text) > 35 or re.match(r"^(describe|explain|provide|list|identify|do|does|is|are|can|will|how|what|which)\b", text, re.I)): continue
        if text.casefold() in seen: continue
        seen.add(text.casefold())
        number = row[number_col].strip() if number_col is not None and number_col < len(row) else ""
        questions.append((number, text))
    return questions

@shared_task(bind=True, autoretry_for=(OSError,), retry_backoff=True, max_retries=3)
def process_questionnaire(self, questionnaire_id):
    questionnaire = Questionnaire.objects.get(id=questionnaire_id)
    questionnaire.processing_status="PROCESSING"; questionnaire.save(update_fields=["processing_status", "updated_at"])
    try:
        parsed = parse_questionnaire(questionnaire)
        if not parsed: raise ValueError("No individual questions could be identified in this file.")
        with transaction.atomic():
            if not questionnaire.questions.exists():
                Question.objects.bulk_create([Question(questionnaire=questionnaire, question_number=number, question_text=text, sort_order=index) for index, (number, text) in enumerate(parsed)], batch_size=500)
            questionnaire.processing_status="COMPLETED"; questionnaire.status="READY"; questionnaire.processing_error=""
            questionnaire.save(update_fields=["processing_status", "status", "processing_error", "updated_at"])
    except Exception as exc:
        questionnaire.processing_status="FAILED"; questionnaire.status="FAILED"; questionnaire.processing_error="Question extraction failed. Review the file format and retry."; questionnaire.save(update_fields=["processing_status", "status", "processing_error", "updated_at"])
        logger.warning("Questionnaire extraction failed: questionnaire_id=%s error_type=%s", questionnaire_id, type(exc).__name__)
        raise
