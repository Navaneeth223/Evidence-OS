import csv
import hashlib
import io
import logging
from celery import shared_task
from django.db import transaction
from django.db.utils import OperationalError
from .models import DocumentVersion, DocumentSection, EvidenceAtom

logger = logging.getLogger(__name__)

def parse_version(version):
    suffix = version.original_filename.lower().rsplit(".", 1)[-1]
    sections = []
    with version.file.open("rb") as raw:
        data = raw.read()
    if suffix == "pdf":
        import fitz
        with fitz.open(stream=data, filetype="pdf") as pdf:
            for page_number, page in enumerate(pdf, 1):
                text = page.get_text("text").strip()
                if text: sections.append(("PAGE", page_number, "", text))
    elif suffix == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text: sections.append(("PARAGRAPH", None, "", text))
        for table_index, table in enumerate(doc.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" | "): sections.append(("TABLE", None, f"Table {table_index}, row {row_index}", text))
    elif suffix == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values: sections.append(("ROW", row_number, sheet.title, " | ".join(values)))
        finally: book.close()
    elif suffix == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        for row_number, row in enumerate(csv.reader(io.StringIO(text)), 1):
            values = [cell.strip() for cell in row if cell.strip()]
            if values: sections.append(("ROW", row_number, "CSV", " | ".join(values)))
    else:
        text = data.decode("utf-8-sig", errors="replace")
        for paragraph in text.splitlines():
            paragraph = paragraph.strip()
            if paragraph: sections.append(("PARAGRAPH", None, "", paragraph))
    return sections

@shared_task(bind=True, autoretry_for=(OSError, OperationalError), retry_backoff=True, max_retries=3)
def process_document(self, version_id):
    version = DocumentVersion.objects.select_related("document", "document__organization").get(id=version_id)
    doc = version.document
    version.processing_status = "PROCESSING"; version.save(update_fields=["processing_status"])
    doc.status = "PROCESSING"; doc.save(update_fields=["status", "updated_at"])
    try:
        parsed = parse_version(version)
        if not parsed: raise ValueError("No readable text was found in this document.")
        with transaction.atomic():
            DocumentSection.objects.filter(version=version).delete()
            EvidenceAtom.objects.filter(version=version).delete()
            atoms = []
            for section_type, location, sheet, content in parsed:
                section = DocumentSection.objects.create(version=version, section_type=section_type, page_number=location if section_type == "PAGE" else None, sheet_name=sheet, section_path=sheet or "", content=content)
                locator = {"document_id": str(doc.id), "version_id": str(version.id), "section_id": str(section.id)}
                if section_type == "PAGE": locator["page"] = location
                elif location is not None: locator["row"] = location
                if sheet: locator["sheet"] = sheet
                digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
                atoms.append(EvidenceAtom(organization=doc.organization, document=doc, version=version, content=content[:12000], title=content[:100], evidence_type="OTHER", source_locator=locator, content_hash=digest))
            EvidenceAtom.objects.bulk_create(atoms, batch_size=500)
            version.processing_status = "COMPLETED"; version.save(update_fields=["processing_status"])
            doc.status = "READY"; doc.save(update_fields=["status", "updated_at"])
    except Exception as exc:
        version.processing_status = "FAILED"; version.save(update_fields=["processing_status"])
        doc.status = "FAILED"; doc.save(update_fields=["status", "updated_at"])
        logger.warning("Document processing failed: version_id=%s error_type=%s", version_id, type(exc).__name__)
        raise
